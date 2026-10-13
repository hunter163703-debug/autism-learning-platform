# -*- coding: utf-8 -*-
"""
自閉症與溝通障礙教學平台
第五節 (Session 1) 與 第六節 (Session 3 & 4) 學生分組學習單 Excel 報表匯出工具
產出檔案：
  1. 第五節_Session1_學生學習單成績總表.xlsx
  2. 第六節_Session3_4_學生學習單成績總表.xlsx
"""

import os
import json
import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SQLITE_DIR = r"D:\2026ant線上教學網站\.wrangler\state\v3\d1\miniflare-D1DatabaseObject"

def get_db_connection():
    if not os.path.exists(SQLITE_DIR):
        return None
    for f in os.listdir(SQLITE_DIR):
        if f.endswith('.sqlite') and not f.startswith('metadata'):
            db_path = os.path.join(SQLITE_DIR, f)
            return sqlite3.connect(db_path)
    return None

def export_session5():
    conn = get_db_connection()
    if not conn:
        print("未找到本地 D1 SQLite 資料庫！")
        return

    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    try:
        cur.execute("SELECT id, student_id, student_name, group_name, group_members, total_score, score_part1, score_part2, score_part3, score_part4, answers_json, ai_feedback, teacher_notes, submitted_at FROM session5_submissions ORDER BY submitted_at DESC")
        rows = cur.fetchall()
    except Exception as e:
        print(f"查詢 session5_submissions 失敗: {e}")
        conn.close()
        return

    data = [dict(r) for r in rows]
    conn.close()

    out_file = r"D:\2026ant線上教學網站\第五節_Session1_學生學習單成績總表.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "最新成績總表"

    headers1 = ["項次", "組別", "學號", "姓名", "同組成員", "總分(100)", "壹、互動起點(25)", "貳、如何與為什麼(35)", "參、跟隨引導OWL(20)", "肆、辨識訊號(20)", "教師課堂註記", "提交時間"]
    ws1.append(headers1)

    # 去重
    seen_students = set()
    unique_rows = []
    for r in data:
        if r['student_id'] not in seen_students:
            seen_students.add(r['student_id'])
            unique_rows.append(r)

    for idx, r in enumerate(unique_rows):
        ws1.append([
            idx + 1,
            r.get('group_name') or '個人',
            r.get('student_id', ''),
            r.get('student_name', ''),
            r.get('group_members', ''),
            r.get('total_score', 0),
            r.get('score_part1', 0),
            r.get('score_part2', 0),
            r.get('score_part3', 0),
            r.get('score_part4', 0),
            r.get('teacher_notes', '') or '',
            r.get('submitted_at', '')
        ])

    ws2 = wb.create_sheet(title="歷次作答詳細清單")
    headers2 = ["紀錄ID", "組別", "學號", "姓名", "總分", "完整作答JSON", "提交時間"]
    ws2.append(headers2)
    for r in data:
        ws2.append([
            r.get('id'),
            r.get('group_name') or '個人',
            r.get('student_id', ''),
            r.get('student_name', ''),
            r.get('total_score', 0),
            r.get('answers_json', ''),
            r.get('submitted_at', '')
        ])

    wb.save(out_file)
    print(f"第五節報表已成功儲存至: {out_file} (共 {len(data)} 筆紀錄，{len(unique_rows)} 位學生)")

def export_session6():
    conn = get_db_connection()
    if not conn:
        print("未找到本地 D1 SQLite 資料庫！")
        return

    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    try:
        cur.execute("SELECT id, student_id, student_name, group_name, group_members, total_score, score_part1, score_part2, answers_json, ai_feedback, teacher_notes, submitted_at FROM session6_submissions ORDER BY submitted_at DESC")
        rows = cur.fetchall()
    except Exception as e:
        print(f"查詢 session6_submissions 失敗: {e}")
        conn.close()
        return

    data = [dict(r) for r in rows]
    conn.close()

    out_file = r"D:\2026ant線上教學網站\第六節_Session3_4_學生學習單成績總表.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "最新成績總表"

    headers1 = ["項次", "組別", "學號", "姓名", "同組成員", "總分(100)", "Part A: Session 3(50)", "Part B: Session 4(50)", "教師課堂註記", "提交時間"]
    ws1.append(headers1)

    seen_students = set()
    unique_rows = []
    for r in data:
        if r['student_id'] not in seen_students:
            seen_students.add(r['student_id'])
            unique_rows.append(r)

    for idx, r in enumerate(unique_rows):
        ws1.append([
            idx + 1,
            r.get('group_name') or '個人',
            r.get('student_id', ''),
            r.get('student_name', ''),
            r.get('group_members', ''),
            r.get('total_score', 0),
            r.get('score_part1', 0),
            r.get('score_part2', 0),
            r.get('teacher_notes', '') or '',
            r.get('submitted_at', '')
        ])

    ws2 = wb.create_sheet(title="歷次作答詳細清單")
    headers2 = ["紀錄ID", "組別", "學號", "姓名", "總分", "完整作答JSON", "提交時間"]
    ws2.append(headers2)
    for r in data:
        ws2.append([
            r.get('id'),
            r.get('group_name') or '個人',
            r.get('student_id', ''),
            r.get('student_name', ''),
            r.get('total_score', 0),
            r.get('answers_json', ''),
            r.get('submitted_at', '')
        ])

    wb.save(out_file)
    print(f"第六節報表已成功儲存至: {out_file} (共 {len(data)} 筆紀錄，{len(unique_rows)} 位學生)")

if __name__ == '__main__':
    print("=== 開始匯出第五節與第六節學習單 Excel 報表 ===")
    export_session5()
    export_session6()
    print("=== 匯出完成 ===")
