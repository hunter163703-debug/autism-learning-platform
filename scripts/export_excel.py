# -*- coding: utf-8 -*-
"""
自閉症與溝通障礙教學平台 - 學生作答成績、錯題深入剖析與 AI 蘇格拉底破障全歷程 Excel 報表匯出工具
產出檔案路徑：d:\2026ant線上教學網站\學生作答成績與錯題報表.xlsx
包含五大專業工作表：
  1. 學生測驗總表 (Student Quiz Overview)
  2. 錯題深度逐項剖析 (Detailed Mistakes Analysis)
  3. 全班難點與題目鑑別分析 (Item Difficulty & Teaching Diagnosis)
  4. AI蘇格拉底破障總覽 (AI Socratic Sessions Summary)
  5. AI對話逐字歷程與思維演進 (AI Socratic Full Transcript & Thought Evolution)
"""

import os
import json
import sqlite3
import urllib.request
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================================
# 1. 資料庫與 API 資料擷取函式
# ==========================================================

def get_submissions():
    """取得所有學生測驗提交紀錄"""
    try:
        req = urllib.request.Request('http://127.0.0.1:8787/api/submissions', headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=3) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            if data.get('success'):
                print(f"成功從 API 取得 {len(data.get('data', []))} 筆學生作答紀錄")
                return data.get('data', [])
    except Exception as e:
        print(f"API 請求學生測驗資料未成功 ({e})，切換讀取本機 SQLite...")

    sqlite_dir = r"d:\2026ant線上教學網站\.wrangler\state\v3\d1\miniflare-D1DatabaseObject"
    if os.path.exists(sqlite_dir):
        for f in os.listdir(sqlite_dir):
            if f.endswith('.sqlite') and not f.startswith('metadata'):
                db_path = os.path.join(sqlite_dir, f)
                try:
                    conn = sqlite3.connect(db_path)
                    conn.row_factory = sqlite3.Row
                    cursor = conn.cursor()
                    cursor.execute("SELECT id, chapter_id, student_id, student_name, score, correct_count, total_questions, wrong_questions, duration_seconds, submitted_at FROM student_submissions ORDER BY id ASC")
                    rows = cursor.fetchall()
                    results = []
                    for r in rows:
                        row_dict = dict(r)
                        try:
                            row_dict['wrong_questions'] = json.loads(row_dict['wrong_questions']) if row_dict['wrong_questions'] else []
                        except:
                            pass
                        results.append(row_dict)
                    conn.close()
                    print(f"成功從本機 SQLite 讀取 {len(results)} 筆測驗紀錄")
                    return results
                except Exception as ex:
                    print(f"讀取 SQLite 測驗紀錄失敗: {ex}")
    return []

def get_tutor_sessions():
    """取得所有 AI 蘇格拉底家教對話會話總覽"""
    try:
        req = urllib.request.Request('http://127.0.0.1:8787/api/ai/tutor/sessions', headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=3) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            if data.get('success'):
                print(f"成功從 API 取得 {len(data.get('data', []))} 筆 AI 家教歷程會話")
                return data.get('data', [])
    except Exception as e:
        print(f"API 請求 AI 會話未成功 ({e})，切換讀取本機 SQLite...")

    sqlite_dir = r"d:\2026ant線上教學網站\.wrangler\state\v3\d1\miniflare-D1DatabaseObject"
    if os.path.exists(sqlite_dir):
        for f in os.listdir(sqlite_dir):
            if f.endswith('.sqlite') and not f.startswith('metadata'):
                db_path = os.path.join(sqlite_dir, f)
                try:
                    conn = sqlite3.connect(db_path)
                    conn.row_factory = sqlite3.Row
                    cursor = conn.cursor()
                    cursor.execute("""
                        SELECT id, student_id, student_name, chapter_id, question_id, 
                               question_tag, question_text, student_choice, correct_choice, 
                               status, turn_count, created_at, updated_at 
                        FROM ai_tutor_sessions ORDER BY id ASC
                    """)
                    rows = cursor.fetchall()
                    results = [dict(r) for r in rows]
                    conn.close()
                    print(f"成功從 SQLite 讀取 {len(results)} 筆 AI 家教歷程會話")
                    return results
                except Exception as ex:
                    print(f"讀取 SQLite 會話失敗: {ex}")
    return []

def get_tutor_messages():
    """取得所有 AI 蘇格拉底家教逐字對話歷程"""
    try:
        req = urllib.request.Request('http://127.0.0.1:8787/api/ai/tutor/all-messages', headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=3) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            if data.get('success'):
                print(f"成功從 API 取得 {len(data.get('data', []))} 則 AI 家教逐字稿紀錄")
                return data.get('data', [])
    except Exception as e:
        print(f"API 請求 AI 逐字紀錄未成功 ({e})，切換讀取本機 SQLite...")

    sqlite_dir = r"d:\2026ant線上教學網站\.wrangler\state\v3\d1\miniflare-D1DatabaseObject"
    if os.path.exists(sqlite_dir):
        for f in os.listdir(sqlite_dir):
            if f.endswith('.sqlite') and not f.startswith('metadata'):
                db_path = os.path.join(sqlite_dir, f)
                try:
                    conn = sqlite3.connect(db_path)
                    conn.row_factory = sqlite3.Row
                    cursor = conn.cursor()
                    cursor.execute("""
                        SELECT 
                            m.id as message_id,
                            m.session_id,
                            m.sender,
                            m.content,
                            m.created_at,
                            s.student_id,
                            s.student_name,
                            s.question_id,
                            s.question_tag,
                            s.question_text,
                            s.student_choice,
                            s.correct_choice,
                            s.status as session_status
                        FROM ai_tutor_messages m
                        LEFT JOIN ai_tutor_sessions s ON m.session_id = s.id
                        ORDER BY m.session_id ASC, m.id ASC
                    """)
                    rows = cursor.fetchall()
                    results = [dict(r) for r in rows]
                    conn.close()
                    print(f"成功從 SQLite 讀取 {len(results)} 則 AI 家教逐字稿紀錄")
                    return results
                except Exception as ex:
                    print(f"讀取 SQLite 逐字紀錄失敗: {ex}")
    return []

# ==========================================================
# 2. 輔助格式化工具
# ==========================================================

def format_duration(seconds):
    try:
        s = int(seconds)
        m = s // 60
        sec = s % 60
        if m == 0:
            return f"{sec} 秒"
        return f"{m} 分 {sec:02d} 秒"
    except:
        return f"{seconds} 秒"

def format_taiwan_time(utc_str):
    if not utc_str:
        return "-"
    try:
        import datetime
        dt_str = str(utc_str).replace('Z', '+00:00')
        if '+' not in dt_str and len(dt_str) == 19:
            dt_str += '+00:00'
        dt = datetime.datetime.fromisoformat(dt_str)
        dt_tw = dt + datetime.timedelta(hours=8)
        return dt_tw.strftime('%Y-%m-%d %H:%M:%S')
    except:
        return str(utc_str)

# 10 題專屬教學診斷與盲點補強建議 (針對教師檢視與教學調整)
TEACHING_DIAGNOSIS = {
    1: "【DSM-IV 亞斯伯格症鑑別】重點在於 AS 在早期語言與認知里程碑上無顯著落後；過動衝動為 ADHD 特徵，非 AS 診斷要件。建議教師多強調「早期語言發展是否遲緩」為 DSM-IV 區分 AS 與典型自閉症之關鍵。",
    2: "【盛行率與神經生物本質】男女比例約 4:1，屬異質性光譜障礙。應加強破除早期「冷漠教養／冰箱母親（Refrigerator Mother）」之歷史偽科學，建立多基因與神經生理機轉認知。",
    3: "【歷史源起與智力分佈】Leo Kanner 於 1943 年提出；傳統文獻顯示自閉症伴隨智能障礙約 50%～70%，遠非 25%。建議教師加強非口語智力測驗在 ASD 評估中的應用說明。",
    4: "【診斷演進歷史定位】美國精神醫學會 DSM-IV 於 1994 年已正式將 AS 列為獨立診斷。建議說明診斷手冊修訂之歷史背景與演進脈絡。",
    5: "【ADHD 衝動過動鑑別】跑跳、搶話、未聽完指令即行動屬於典型執行功能與抑制控制缺陷（ADHD）。建議教師指導學生鑑別「語用衝動」與「自閉症心智理論困難」之本質差異。",
    6: "【AS 早期里程碑】AS 具備社交與刻板行為，但無臨床顯著之一般語言遲緩（2歲具單字、3歲具片語）。需與語言發展遲緩（DLD）與典型 ASD 明確區分。",
    7: "【DLD 與 ASD 語用差異】DLD 之語用困難次發於詞彙語法工具不足；ASD 則為社交心智理論（ToM）與情緒互惠之本質障礙。此為臨床語病鑑別之極高頻混淆點，應加強重點講解。",
    8: "【DSM-5 診斷架構演進】DSM-5 將口語遲緩排除在必要核心之外（改為附加註記），統整為社交溝通與侷限行為二合一向度。需向學生講授為何排除口語遲緩有助於光譜整合。",
    9: "【神經解剖與致病迷思】自閉症涉及多基因遺傳及杏仁核（Amygdala）神經發育異常；國際醫學大規模證實與 MMR 疫苗或父母教養風格毫無因果關係。",
    10: "【仿說（Echolalia）之溝通功能】鸚鵡式仿說並非無效的刻板行為，而具有維持互動、請求、抗議與資訊處理調節功能。應引導學生以優勢觀點介入，將仿說轉化為實質溝通橋樑。"
}

# ==========================================================
# 3. 主匯出程式
# ==========================================================

def export_to_excel():
    submissions = get_submissions()
    tutor_sessions = get_tutor_sessions()
    tutor_messages = get_tutor_messages()

    print(f"彙整完畢：{len(submissions)} 筆測驗提交, {len(tutor_sessions)} 筆家教對話, {len(tutor_messages)} 則對話逐字紀錄")

    wb = Workbook()

    # 字型樣式
    header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="微軟正黑體", size=13, bold=True, color="1E293B")
    body_font = Font(name="微軟正黑體", size=10)
    bold_body_font = Font(name="微軟正黑體", size=10, bold=True)
    mono_font = Font(name="Consolas", size=9.5)

    # 填色定義 (專業多主題色系)
    fill_blue_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")    # 深藍 (總表)
    fill_sky_header = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")     # 天藍 (錯題剖析)
    fill_green_header = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")   # 墨綠 (難題統計)
    fill_purple_header = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")  # 深紫 (AI總覽)
    fill_indigo_header = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")  # 靛青 (AI逐字稿)

    # 行級強調底色
    fill_student_msg = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")    # 淡藍 (學生發言)
    fill_tutor_msg = PatternFill(start_color="FAF5FF", end_color="FAF5FF", fill_type="solid")      # 淡紫 (AI家教發言)
    fill_clarified = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")      # 淡綠 (突破成功)
    fill_warning = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")        # 淡黃 (進行中/難題)

    # 框線定義
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # -------------------------------------------------------------------------
    # 工作表 1：學生測驗總表 (Student Quiz Overview)
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "學生測驗總表"
    ws1.views.sheetView[0].showGridLines = True

    headers1 = [
        "記錄序號", "提交時間 (台灣時間)", "課程章節", "學號", "姓名",
        "測驗分數", "答對題數", "總題數", "正確率", "作答秒數", "作答時長",
        "錯題數", "需加強錯題題號列表"
    ]
    ws1.append(headers1)
    ws1.row_dimensions[1].height = 28

    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill_blue_header
        cell.alignment = align_center
        cell.border = thin_border

    row_num1 = 2
    for s in submissions:
        wrong_list = s.get('wrong_questions', [])
        if not isinstance(wrong_list, list):
            wrong_list = []
        wrong_orders = ", ".join([f"第{w.get('sort_order')}題" for w in wrong_list]) if wrong_list else "全對 (0題)"
        accuracy = f"{round((s.get('correct_count', 0) / (s.get('total_questions') or 10)) * 100)}%"

        row_data1 = [
            s.get('id', ''),
            format_taiwan_time(s.get('submitted_at', '')),
            "第一節 自閉症兒童的診斷以及臨床表現",
            s.get('student_id', ''),
            s.get('student_name', ''),
            s.get('score', 0),
            s.get('correct_count', 0),
            s.get('total_questions', 10),
            accuracy,
            s.get('duration_seconds', 0),
            format_duration(s.get('duration_seconds', 0)),
            len(wrong_list),
            wrong_orders
        ]
        ws1.append(row_data1)
        score_val = s.get('score', 0)

        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=row_num1, column=col_idx)
            cell.font = body_font
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6, 7, 8, 9, 10, 11, 12]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            # 分數色彩反饋
            if col_idx == 6:
                cell.font = bold_body_font
                if score_val >= 80:
                    cell.fill = fill_clarified
                elif score_val < 60:
                    cell.fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")

        ws1.row_dimensions[row_num1].height = 22
        row_num1 += 1

    if not submissions:
        ws1.append(["-", "尚無作答提交紀錄", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-"])
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=2, column=col_idx)
            cell.font = body_font
            cell.alignment = align_center
            cell.border = thin_border

    # -------------------------------------------------------------------------
    # 工作表 2：錯題深度逐項剖析 (Detailed Mistakes Analysis)
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="錯題深度剖析")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = [
        "記錄序號", "提交時間", "學號", "姓名", "錯題題號", "概念分類標籤",
        "題目完整內容", "學生所選選項", "學生選項錯誤解析",
        "標準解答選項", "標準解答深度剖析", "核心概念歸納與反思"
    ]
    ws2.append(headers2)
    ws2.row_dimensions[1].height = 28

    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill_sky_header
        cell.alignment = align_center
        cell.border = thin_border

    row_num2 = 2
    for s in submissions:
        w_list = s.get('wrong_questions', [])
        if isinstance(w_list, list):
            for w in w_list:
                row_data2 = [
                    s.get('id', ''),
                    format_taiwan_time(s.get('submitted_at', '')),
                    s.get('student_id', ''),
                    s.get('student_name', ''),
                    f"第 {w.get('sort_order')} 題",
                    w.get('source_tag', ''),
                    w.get('question_text', ''),
                    f"【{w.get('selected_key')}】{w.get('selected_text', '')}",
                    w.get('selected_explanation', ''),
                    f"【{w.get('correct_key')}】{w.get('correct_text', '')}",
                    w.get('correct_explanation', ''),
                    w.get('summary', '')
                ]
                ws2.append(row_data2)
                for col_idx in range(1, len(headers2) + 1):
                    cell = ws2.cell(row=row_num2, column=col_idx)
                    cell.font = body_font
                    cell.border = thin_border
                    if col_idx in [1, 2, 3, 4, 5]:
                        cell.alignment = align_center
                    elif col_idx in [8, 10]:
                        cell.font = bold_body_font
                        cell.alignment = align_left
                    else:
                        cell.alignment = align_left
                ws2.row_dimensions[row_num2].height = 36
                row_num2 += 1

    if row_num2 == 2:
        ws2.append(["-", "尚無錯題紀錄 (全體滿分或尚未作答)", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-"])
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=2, column=col_idx)
            cell.font = body_font
            cell.alignment = align_center
            cell.border = thin_border

    # -------------------------------------------------------------------------
    # 工作表 3：全班難點與題目鑑別分析 (Item Difficulty & Teaching Diagnosis)
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="全班難點與教學鑑別")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = [
        "題號", "概念主題標籤", "標準正解", "受測總人次", "全班答錯人次",
        "答錯率 (難度指標)", "錯誤選項分佈詳情", "教學盲點診斷與後續授課補強建議"
    ]
    ws3.append(headers3)
    ws3.row_dimensions[1].height = 28

    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill_green_header
        cell.alignment = align_center
        cell.border = thin_border

    sec1_path = r"d:\2026ant線上教學網站\content\section-1.json"
    questions_meta = []
    if os.path.exists(sec1_path):
        with open(sec1_path, 'r', encoding='utf-8') as f:
            sec1_data = json.load(f)
            questions_meta = sec1_data.get('questions', [])

    total_submissions_count = len(submissions)
    for idx in range(1, 11):
        meta = questions_meta[idx - 1] if idx <= len(questions_meta) else {}
        tag = meta.get('source', '')
        correct_ans = meta.get('correct', '')

        wrong_count_q = 0
        picked_keys = {}
        for s in submissions:
            w_list = s.get('wrong_questions', [])
            if isinstance(w_list, list):
                for w in w_list:
                    if w.get('sort_order') == idx:
                        wrong_count_q += 1
                        k = w.get('selected_key', '')
                        picked_keys[k] = picked_keys.get(k, 0) + 1

        wrong_ratio = (wrong_count_q / total_submissions_count) if total_submissions_count > 0 else 0
        wrong_rate_str = f"{(wrong_ratio * 100):.1f}%"
        
        if total_submissions_count == 0:
            picked_summary = "尚無測驗數據"
        else:
            picked_summary = ", ".join([f"選【{k}】: {v} 人" for k, v in sorted(picked_keys.items())]) if picked_keys else "全對無錯誤"

        diagnosis = TEACHING_DIAGNOSIS.get(idx, "請強化本題核心觀念之臨床實例演練。")

        row_data3 = [
            f"第 {idx} 題",
            tag,
            f"【{correct_ans}】",
            total_submissions_count,
            wrong_count_q,
            wrong_rate_str,
            picked_summary,
            diagnosis
        ]
        ws3.append(row_data3)
        row_num3 = idx + 1

        for col_idx in range(1, len(headers3) + 1):
            cell = ws3.cell(row=row_num3, column=col_idx)
            cell.font = body_font
            cell.border = thin_border
            if col_idx in [1, 2, 3, 4, 5, 6]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            # 若答錯率高於 30%，以醒目色彩標註以利教師立即辨識
            if col_idx == 6 and wrong_ratio >= 0.3:
                cell.fill = fill_warning
                cell.font = bold_body_font

        ws3.row_dimensions[row_num3].height = 42

    # -------------------------------------------------------------------------
    # 工作表 4：AI蘇格拉底破障總覽 (AI Socratic Sessions Summary)
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="AI家教破障總覽")
    ws4.views.sheetView[0].showGridLines = True

    headers4 = [
        "歷程序號", "對話 ID (Session)", "啟動時間", "學號", "姓名",
        "檢測題號", "概念主題標籤", "題目簡述", "學生初選錯誤",
        "標準解答", "對話輪數", "破障狀態", "最後更新時間"
    ]
    ws4.append(headers4)
    ws4.row_dimensions[1].height = 28

    for col_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill_purple_header
        cell.alignment = align_center
        cell.border = thin_border

    row_num4 = 2
    for idx, s in enumerate(tutor_sessions, 1):
        is_clarified = (s.get('status') == 'clarified')
        status_text = "🎉 觀念已突破" if is_clarified else "💡 引導進行中"
        sid = s.get('id') or s.get('session_id') or ''
        qid = s.get('question_id') or s.get('question_sort_order') or 1
        turn_num = s.get('turn_count') or s.get('turns_count') or 0
        s_choice = s.get('student_choice') or s.get('selected_key') or ''
        c_choice = s.get('correct_choice') or s.get('correct_key') or ''

        row_data4 = [
            idx,
            f"SESSION-{sid}",
            format_taiwan_time(s.get('created_at', '')),
            s.get('student_id', ''),
            s.get('student_name', ''),
            f"第 {qid} 題",
            s.get('question_tag', '核心觀念'),
            (s.get('question_text', '') or '')[:40] + '...',
            f"【{s_choice}】",
            f"【{c_choice}】",
            turn_num,
            status_text,
            format_taiwan_time(s.get('updated_at', ''))
        ]
        ws4.append(row_data4)

        for col_idx in range(1, len(headers4) + 1):
            cell = ws4.cell(row=row_num4, column=col_idx)
            cell.font = body_font
            cell.border = thin_border
            if col_idx in [1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if col_idx == 12:
                cell.font = bold_body_font
                cell.fill = fill_clarified if is_clarified else fill_warning

        ws4.row_dimensions[row_num4].height = 24
        row_num4 += 1

    if not tutor_sessions:
        ws4.append(["-", "尚無 AI 家教對話紀錄", "-", "-", "-", "-", "-", "-", "-", "-", 0, "-", "-"])
        for col_idx in range(1, len(headers4) + 1):
            cell = ws4.cell(row=2, column=col_idx)
            cell.font = body_font
            cell.alignment = align_center
            cell.border = thin_border

    # -------------------------------------------------------------------------
    # 工作表 5：AI對話逐字稿與思維演進歷程 (AI Socratic Full Transcript)
    # -------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="AI逐字對話歷程")
    ws5.views.sheetView[0].showGridLines = True

    headers5 = [
        "訊息序號", "對話會話 ID", "學號", "姓名", "探討題號",
        "發言者角色", "完整對話發言與引導提問內容", "思維演進與破障階段註記", "發言時間 (台灣時間)"
    ]
    ws5.append(headers5)
    ws5.row_dimensions[1].height = 28

    for col_idx, h in enumerate(headers5, 1):
        cell = ws5.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill_indigo_header
        cell.alignment = align_center
        cell.border = thin_border

    row_num5 = 2
    for idx, m in enumerate(tutor_messages, 1):
        sender = m.get('sender', '')
        is_student = (sender == 'student' or sender == 'user')
        sender_label = "👤 學生提問／回應" if is_student else "🤖 AI 蘇格拉底家教"

        content = m.get('content', '') or ''
        
        # 判定思維階段
        stage_note = "初探思維卡點與盲點釐清"
        if is_student:
            stage_note = "學生表達個人認知推論與困惑"
        else:
            if "恭喜" in content or "完全正確" in content or "太棒了" in content:
                stage_note = "🎉 觀念成功突破！融會貫通"
            elif idx > 2:
                stage_note = "階梯式反詰問深化思考"

        sid = m.get('session_id', '')
        qid = m.get('question_id', '')

        row_data5 = [
            idx,
            f"SESSION-{sid}",
            m.get('student_id', ''),
            m.get('student_name', ''),
            f"第 {qid} 題" if qid else "-",
            sender_label,
            content,
            stage_note,
            format_taiwan_time(m.get('created_at', ''))
        ]
        ws5.append(row_data5)

        for col_idx in range(1, len(headers5) + 1):
            cell = ws5.cell(row=row_num5, column=col_idx)
            cell.font = body_font
            cell.border = thin_border
            
            # 發言色彩區隔
            if is_student:
                cell.fill = fill_student_msg
            else:
                cell.fill = fill_tutor_msg

            if col_idx in [1, 2, 3, 4, 5, 6, 8, 9]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            if col_idx == 6:
                cell.font = bold_body_font

            if col_idx == 8 and "突破" in stage_note:
                cell.fill = fill_clarified
                cell.font = bold_body_font

        # 動態調整列高以容納長對話
        content_lines = len(content.split('\n'))
        approx_lines = max(content_lines, len(content) // 60 + 1)
        ws5.row_dimensions[row_num5].height = max(32, min(approx_lines * 18, 140))
        row_num5 += 1

    if not tutor_messages:
        ws5.append(["-", "尚無 AI 對話逐字紀錄", "-", "-", "-", "-", "尚無對話", "-", "-"])
        for col_idx in range(1, len(headers5) + 1):
            cell = ws5.cell(row=2, column=col_idx)
            cell.font = body_font
            cell.alignment = align_center
            cell.border = thin_border

    # -------------------------------------------------------------------------
    # 欄寬自動適應與手動關鍵欄位精緻化微調
    # -------------------------------------------------------------------------
    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                # 中文字計長 2，英數計長 1
                length = sum(2 if ord(char) > 127 else 1 for char in val.split('\n')[0])
                if length > max_len:
                    max_len = length
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 12)

    # 關鍵欄位精緻微調
    ws1.column_dimensions['B'].width = 22  # 時間
    ws1.column_dimensions['C'].width = 38  # 章節名稱
    ws1.column_dimensions['M'].width = 30  # 錯題列表

    ws2.column_dimensions['F'].width = 30  # 標籤
    ws2.column_dimensions['G'].width = 42  # 題目
    ws2.column_dimensions['H'].width = 38  # 學生選項
    ws2.column_dimensions['I'].width = 45  # 錯誤解析
    ws2.column_dimensions['J'].width = 38  # 正解選項
    ws2.column_dimensions['K'].width = 45  # 正解剖析
    ws2.column_dimensions['L'].width = 40  # 核心歸納

    ws3.column_dimensions['B'].width = 32  # 主題標籤
    ws3.column_dimensions['G'].width = 35  # 選項分佈
    ws3.column_dimensions['H'].width = 65  # 教學診斷建議

    ws4.column_dimensions['C'].width = 22  # 時間
    ws4.column_dimensions['G'].width = 30  # 主題
    ws4.column_dimensions['H'].width = 45  # 題目簡述
    ws4.column_dimensions['M'].width = 22  # 更新時間

    ws5.column_dimensions['B'].width = 18  # Session ID
    ws5.column_dimensions['F'].width = 20  # 發言角色
    ws5.column_dimensions['G'].width = 75  # 對話內容 (加寬以便閱讀全文)
    ws5.column_dimensions['H'].width = 32  # 階段註記
    ws5.column_dimensions['I'].width = 22  # 發言時間

    # -------------------------------------------------------------------------
    # 存檔與產出回報
    # -------------------------------------------------------------------------
    target_excel = r"d:\2026ant線上教學網站\學生作答成績與錯題報表.xlsx"
    target_csv = r"d:\2026ant線上教學網站\學生作答成績總表.csv"

    wb.save(target_excel)
    print(f"🎉 5 大工作表專業 Excel 報表已成功產出至: {target_excel}")

    # 產出總表 CSV 備查
    try:
        df1 = pd.read_excel(target_excel, sheet_name="學生測驗總表")
        df1.to_csv(target_csv, index=False, encoding='utf-8-sig')
        print(f"✅ CSV 總表已成功產出至: {target_csv}")
    except Exception as e:
        print(f"產出 CSV 提示: {e}")

if __name__ == '__main__':
    export_to_excel()